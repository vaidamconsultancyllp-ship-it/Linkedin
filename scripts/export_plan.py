"""Export data/plan.json to a branded Excel sheet and PDF in plan/.

    python3 scripts/export_plan.py

Needs openpyxl and reportlab (pip install openpyxl reportlab).
"""
import json
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from common import DATA, ROOT

NAVY, TEAL, BLUE, GOLD, ORANGE = "021422", "028682", "008FB5", "EFAA30", "EB642D"
PILLAR_COLORS = {
    "Deadline alerts": "FDE3D6",
    "Rule changes": "D6EEF6",
    "Explainers": "D5EEED",
    "Costly mistakes": "FBEBCD",
    "Founder checklists": "E4E8EC",
    "Founder FAQ": "E9F5F4",
    "Myth vs fact": "FFF4E0",
}
HEADERS = ["Date", "Day", "Theme", "Format", "Topic", "Key points", "Status"]


def register_fonts():
    """Use a TTF font with the rupee sign; the built-in Helvetica lacks it."""
    candidates = [
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        ("/Library/Fonts/Arial Unicode.ttf", "/Library/Fonts/Arial Unicode.ttf"),
    ]
    for regular, bold in candidates:
        try:
            pdfmetrics.registerFont(TTFont("Body", regular))
            pdfmetrics.registerFont(TTFont("Body-Bold", bold))
            return "Body", "Body-Bold"
        except Exception:
            continue
    return "Helvetica", "Helvetica-Bold"


def fmt_date(iso):
    return date.fromisoformat(iso).strftime("%d %b %Y")


def export_xlsx(plan, path, span):
    wb = Workbook()
    ws = wb.active
    ws.title = "60-Day Plan"
    arial = "Arial"
    ws["A1"] = "Vaidam Consultancy LLP – LinkedIn Content Plan"
    ws["A1"].font = Font(name=arial, size=14, bold=True, color=NAVY)
    ws["A2"] = f"{span} · one post daily at 10:00 AM IST · topics are re-verified on the day"
    ws["A2"].font = Font(name=arial, size=10, italic=True, color="555555")

    header_row = 4
    thin = Side(style="thin", color="C8D0D8")
    for col, name in enumerate(HEADERS, 1):
        cell = ws.cell(row=header_row, column=col, value=name)
        cell.font = Font(name=arial, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center")
    for i, p in enumerate(plan, header_row + 1):
        values = [date.fromisoformat(p["date"]), p["day"], p["pillar"], p["format"].capitalize(), p["topic"],
                  p["key_points"], "Planned"]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.font = Font(name=arial, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)
        ws.cell(row=i, column=1).number_format = "DD MMM YYYY"
        ws.cell(row=i, column=3).fill = PatternFill("solid", fgColor=PILLAR_COLORS.get(p["pillar"], "FFFFFF"))
    for col, width in zip("ABCDEFG", (13, 6, 18, 12, 60, 55, 10)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    last = header_row + len(plan)
    ws.auto_filter.ref = f"A{header_row}:G{last}"
    ws.cell(row=last + 2, column=1, value="Edit the Status column (Planned / Posted / Skipped) to track progress.").font = \
        Font(name=arial, size=9, italic=True, color="555555")

    wb.save(path)


def export_pdf(plan, path, span):
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            title="Vaidam LinkedIn Content Plan", author="Vaidam Consultancy LLP")
    body, bold = register_fonts()
    title = ParagraphStyle("t", fontName=bold, fontSize=16, textColor=colors.HexColor("#" + NAVY))
    sub = ParagraphStyle("s", fontName=body, fontSize=9, textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("c", fontName=body, fontSize=8, leading=10)
    head = ParagraphStyle("h", fontName=bold, fontSize=8.5, textColor=colors.white)

    rows = [[Paragraph(h, head) for h in HEADERS[:6]]]
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, colors.HexColor("#C8D0D8")),
        ("LINEABOVE", (0, 0), (-1, 0), 3, colors.HexColor("#" + TEAL)),
    ]
    for i, p in enumerate(plan, 1):
        rows.append([Paragraph(fmt_date(p["date"]), cell), Paragraph(p["day"], cell), Paragraph(p["pillar"], cell),
                     Paragraph(p["format"].capitalize(), cell),
                     Paragraph(p["topic"].replace("&", "&amp;"), cell),
                     Paragraph(p["key_points"].replace("&", "&amp;"), cell)])
        style.append(("BACKGROUND", (2, i), (2, i), colors.HexColor("#" + PILLAR_COLORS.get(p["pillar"], "FFFFFF"))))
    table = Table(rows, colWidths=[25 * mm, 12 * mm, 29 * mm, 23 * mm, 100 * mm, 84 * mm], repeatRows=1)
    table.setStyle(TableStyle(style))
    story = [Paragraph("Vaidam Consultancy LLP – LinkedIn Content Plan", title), Spacer(1, 3 * mm),
             Paragraph(f"{span} · one post daily at 10:00 AM IST · Mon deadlines · Tue rule changes · "
                       "Wed explainers · Thu costly mistakes · Fri checklists · Sat founder FAQ · Sun myth vs fact. "
                       "Wednesday and Friday posts are one-page infographics; other days are carousels. "
                       "Dates and amounts are re-verified on the day of posting.", sub),
             Spacer(1, 5 * mm), table]
    doc.build(story)


def main():
    plan = json.loads((DATA / "plan.json").read_text())
    span = f"{fmt_date(plan[0]['date'])} – {fmt_date(plan[-1]['date'])}"
    out = ROOT / "plan"
    out.mkdir(exist_ok=True)
    stem = f"content-plan-{plan[0]['date']}-to-{plan[-1]['date']}"
    export_xlsx(plan, out / f"{stem}.xlsx", span)
    export_pdf(plan, out / f"{stem}.pdf", span)
    print(out / f"{stem}.xlsx")
    print(out / f"{stem}.pdf")


if __name__ == "__main__":
    main()
